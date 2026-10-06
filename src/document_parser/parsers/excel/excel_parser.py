import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ...config import ParserConfig
from ...errors import DocumentParseError
from ...normalization.normalizer import Normalizer
from ...normalization.schema import CanonicalDocument
from ...utils.file_utils import document_data, table_markdown


def contiguous_bands(indices: set[int]) -> list[tuple[int, int]]:
    bands = []
    for n in sorted(indices):
        if bands and n == bands[-1][1] + 1:
            bands[-1] = (bands[-1][0], n)
        else:
            bands.append((n, n))
    return bands


def detect_table_regions(sheet) -> list[tuple[int, int, int, int]]:
    """Split populated row bands, then populated column bands within each band.

    Merged rectangles count as occupied to prevent splitting merged titles.
    Formatting-only cells do not influence region boundaries.
    """
    occupied = {(c.row, c.column) for row in sheet.iter_rows() for c in row if c.value is not None}
    for merged in sheet.merged_cells.ranges:
        if sheet.cell(merged.min_row, merged.min_col).value is not None:
            occupied.update((r, c) for r in range(merged.min_row, merged.max_row + 1)
                            for c in range(merged.min_col, merged.max_col + 1))
    regions = []
    for r1, r2 in contiguous_bands({r for r, _ in occupied}):
        columns = {c for r, c in occupied if r1 <= r <= r2}
        for c1, c2 in contiguous_bands(columns):
            subset = [(r, c) for r, c in occupied if r1 <= r <= r2 and c1 <= c <= c2]
            regions.append((min(r for r, _ in subset), c1, max(r for r, _ in subset), c2))
    return sorted(regions)


def json_value(value):
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


class ExcelParser:
    def __init__(self, config: ParserConfig | None = None):
        self.config = config or ParserConfig()

    def parse(self, path: Path) -> CanonicalDocument:
        data = document_data(path, "openpyxl")
        formulas = load_workbook(path, data_only=False)
        cached = load_workbook(path, data_only=True)
        try:
            data["metadata"]["sheets"] = []
            for sheet in formulas:
                # Protect against sheets inflated by formatting to Excel limits.
                if sheet.max_row * sheet.max_column > self.config.max_excel_region_cells:
                    raise DocumentParseError(f"Sheet {sheet.title!r} exceeds configured cell limit")
                sheet_meta = {
                    "sheet": sheet.title, "state": sheet.sheet_state,
                    "merged_cells": [str(r) for r in sheet.merged_cells.ranges],
                    "hidden_rows": [n for n, d in sheet.row_dimensions.items() if d.hidden],
                    "hidden_columns": [k for k, d in sheet.column_dimensions.items() if d.hidden],
                }
                data["metadata"]["sheets"].append(sheet_meta)
                for r1, c1, r2, c2 in detect_table_regions(sheet):
                    rows, cells = [], []
                    for row in sheet.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                        values = []
                        for cell in row:
                            value = cached[sheet.title].cell(cell.row, cell.column).value
                            formula = str(cell.value) if cell.data_type == "f" else None
                            if not formula:
                                value = cell.value
                            values.append(json_value(value) if value is not None else formula)
                            if cell.value is not None:
                                cells.append({"coordinate": cell.coordinate, "row": cell.row, "column": cell.column,
                                              "value": json_value(value), "formula": formula, "data_type": cell.data_type,
                                              "number_format": cell.number_format})
                        rows.append(values)
                    region = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
                    data["elements"].append({"element_type": "table", "text": table_markdown(rows),
                        "metadata": {**sheet_meta, "parser": "openpyxl", "workbook": path.name,
                                     "range": region, "values": rows, "cells": cells,
                                     "formula_cache_note": "openpyxl does not recalculate formulas"}})
        finally:
            formulas.close()
            cached.close()
        return Normalizer().normalize(data)
